"""Excel (.xlsx) reporting export for confirmed bookings.

For hospital staff who live in spreadsheets, every confirmed booking is also
appended as a row to an Excel workbook. This is a *reporting* mirror, not the
source of truth: Google Calendar owns the schedule and SQLite owns the
application data. Consequently a failure here (e.g. the file is open in Excel and
locked) is logged and swallowed — it must never block a booking that has already
been written to the calendar and the database.

The workbook and its header row are created automatically on first append.
"""

from __future__ import annotations

from pathlib import Path
from threading import Lock
from typing import TYPE_CHECKING

from titan.config import Settings, get_settings
from titan.utils import get_logger

if TYPE_CHECKING:  # avoid a runtime import cycle (booking imports this module)
    from titan.services.booking import Booking

logger = get_logger("titan.reporting")

_HEADERS = [
    "Name",
    "Phone",
    "Age",
    "Gender",
    "Department",
    "Doctor",
    "Symptoms",
    "Appointment Date & Time",
    "Google Calendar Event ID",
    "Booking Status",
    "Timestamp",
]


class ExcelReport:
    """Appends confirmed bookings to a staff-facing ``.xlsx`` workbook."""

    def __init__(self, settings: Settings | None = None) -> None:
        self._settings = settings or get_settings()
        self._path: Path = self._settings.excel_path
        self._lock = Lock()

    def append(self, booking: Booking) -> bool:
        """Append ``booking`` as a row. Returns False on failure (never raises)."""
        try:
            from openpyxl import Workbook, load_workbook

            with self._lock:
                self._path.parent.mkdir(parents=True, exist_ok=True)
                if self._path.exists():
                    workbook = load_workbook(self._path)
                    sheet = workbook.active
                else:
                    workbook = Workbook()
                    sheet = workbook.active
                    sheet.title = "Appointments"
                    sheet.append(_HEADERS)

                sheet.append(
                    [
                        booking.name,
                        booking.phone,
                        booking.age,
                        booking.gender,
                        booking.department,
                        booking.doctor,
                        booking.symptoms,
                        booking.start.strftime("%Y-%m-%d %H:%M"),
                        booking.event_id,
                        booking.status,
                        booking.created_at.strftime("%Y-%m-%d %H:%M:%S")
                        if booking.created_at
                        else "",
                    ]
                )
                workbook.save(self._path)
        except Exception as exc:  # reporting must never break a booking
            logger.warning("Excel report append failed (non-fatal): %s", exc)
            return False

        logger.info(
            "Appended booking for %s to Excel report %s", booking.name, self._path
        )
        return True


_report: ExcelReport | None = None


def get_report() -> ExcelReport:
    """Return the process-wide :class:`ExcelReport` singleton."""
    global _report
    if _report is None:
        _report = ExcelReport()
    return _report
