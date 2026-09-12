"""Offline unit tests for the SQLite datastore and Excel report.

Both use temporary paths so they auto-create their files without touching the
real ``data/`` directory, and neither needs network or model access.
"""

from datetime import datetime
from zoneinfo import ZoneInfo

from titan.config import Settings
from titan.services.appointment_db import AppointmentDB
from titan.services.booking import Booking
from titan.services.reporting import ExcelReport

_TZ = ZoneInfo("Asia/Kolkata")


def _booking() -> Booking:
    return Booking(
        name="Asha",
        phone="9876543210",
        age=30,
        gender="Female",
        department="Cardiology",
        doctor="Dr. Rajesh",
        symptoms="chest pain",
        start=datetime(2026, 7, 8, 15, 0, tzinfo=_TZ),
        duration_minutes=30,
        event_id="mem-1",
        status="confirmed",
        created_at=datetime(2026, 7, 5, 10, 0, tzinfo=_TZ),
    )


def test_sqlite_autocreates_and_roundtrips(tmp_path) -> None:
    db = AppointmentDB(Settings(database_file=str(tmp_path / "titan.db")))
    row_id = db.save(_booking())

    assert row_id == 1
    assert (tmp_path / "titan.db").exists()

    rows = db.all()
    assert len(rows) == 1
    row = rows[0]
    assert row["name"] == "Asha"
    assert row["phone"] == "9876543210"
    assert row["age"] == 30
    assert row["gender"] == "Female"
    assert row["department"] == "Cardiology"
    assert row["doctor"] == "Dr. Rajesh"
    assert row["symptoms"] == "chest pain"
    assert row["event_id"] == "mem-1"
    assert row["status"] == "confirmed"
    assert row["appointment_start"].startswith("2026-07-08T15:00")
    assert row["created_at"].startswith("2026-07-05T10:00")


def test_excel_autocreates_header_and_appends(tmp_path) -> None:
    from openpyxl import load_workbook

    path = tmp_path / "appointments.xlsx"
    report = ExcelReport(Settings(excel_file=str(path)))

    assert report.append(_booking()) is True
    assert path.exists()

    sheet = load_workbook(path).active
    assert sheet.cell(row=1, column=1).value == "Name"
    assert sheet.cell(row=1, column=9).value == "Google Calendar Event ID"
    assert sheet.cell(row=2, column=1).value == "Asha"
    assert sheet.cell(row=2, column=9).value == "mem-1"
    assert sheet.max_row == 2  # header + one booking


def test_excel_appends_accumulate(tmp_path) -> None:
    from openpyxl import load_workbook

    path = tmp_path / "appointments.xlsx"
    report = ExcelReport(Settings(excel_file=str(path)))
    report.append(_booking())
    report.append(_booking())

    sheet = load_workbook(path).active
    assert sheet.max_row == 3  # header + two bookings
