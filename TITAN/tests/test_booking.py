"""Offline unit tests for the booking orchestrator.

Wires an in-memory calendar, a fixed directory, and temporary SQLite/Excel files
so the full flow (validate -> schedule -> persist) is exercised without network,
credentials, or a live model.
"""

import pytest

from titan.config import Settings
from titan.services import booking as booking_mod
from titan.services.appointment_db import AppointmentDB
from titan.services.booking import SlotUnavailableError, create_booking
from titan.services.calendar import AppointmentError, InMemoryCalendar
from titan.services.directory import Department, Directory, Doctor
from titan.services.reporting import ExcelReport


@pytest.fixture
def wired(tmp_path, monkeypatch):
    """Point the booking module's seams at test doubles / temp files."""
    settings = Settings(
        business_timezone="Asia/Kolkata",
        business_hours_start=9,
        business_hours_end=18,
        default_appointment_minutes=30,
        alternative_slot_count=3,
        alternative_search_window_minutes=180,
        database_file=str(tmp_path / "titan.db"),
        excel_file=str(tmp_path / "appointments.xlsx"),
    )
    calendar = InMemoryCalendar(settings)
    directory = Directory(
        [
            Department(
                name="Cardiology",
                doctors=[Doctor(name="Dr. Rajesh"), Doctor(name="Dr. Anitha")],
            )
        ]
    )
    db = AppointmentDB(settings)
    report = ExcelReport(settings)

    monkeypatch.setattr(booking_mod, "get_settings", lambda: settings)
    monkeypatch.setattr(booking_mod, "get_calendar", lambda: calendar)
    monkeypatch.setattr(booking_mod, "get_directory", lambda s=None: directory)
    monkeypatch.setattr(booking_mod, "get_appointment_db", lambda: db)
    monkeypatch.setattr(booking_mod, "get_report", lambda: report)
    return settings, calendar, directory, db, report


def test_valid_booking_persists_to_sqlite_and_excel(wired) -> None:
    from openpyxl import load_workbook

    settings, _cal, _dir, db, _report = wired
    booking = create_booking(
        name="Asha",
        department="cardiology",  # case-insensitive
        doctor="Rajesh",  # partial match -> Dr. Rajesh
        preferred_date="2026-07-08",
        preferred_time="15:00",
        phone="9876543210",
        age=30,
        gender="Female",
        symptoms="chest pain",
    )

    assert booking.status == "confirmed"
    assert booking.event_id is not None
    assert booking.department == "Cardiology"
    assert booking.doctor == "Dr. Rajesh"

    rows = db.all()
    assert len(rows) == 1
    assert rows[0]["event_id"] == booking.event_id
    assert rows[0]["age"] == 30

    sheet = load_workbook(settings.excel_path).active
    assert sheet.max_row == 2  # header + booking


def test_unknown_department_raises(wired) -> None:
    with pytest.raises(AppointmentError):
        create_booking(
            name="X",
            department="Oncology",
            doctor="Dr. Rajesh",
            preferred_date="2026-07-08",
            preferred_time="15:00",
        )


def test_unknown_doctor_raises(wired) -> None:
    with pytest.raises(AppointmentError):
        create_booking(
            name="X",
            department="Cardiology",
            doctor="Dr. Nobody",
            preferred_date="2026-07-08",
            preferred_time="15:00",
        )


def test_taken_slot_offers_three_alternatives(wired) -> None:
    _settings, _cal, _dir, db, _report = wired
    create_booking(
        name="First",
        department="Cardiology",
        doctor="Dr. Rajesh",
        preferred_date="2026-07-08",
        preferred_time="15:00",
    )

    with pytest.raises(SlotUnavailableError) as excinfo:
        create_booking(
            name="Second",
            department="Cardiology",
            doctor="Dr. Anitha",
            preferred_date="2026-07-08",
            preferred_time="15:00",
        )

    alternatives = excinfo.value.alternatives
    assert len(alternatives) == 3
    # The clashing 15:00 slot must not be offered back.
    assert all(alt.strftime("%H:%M") != "15:00" for alt in alternatives)
    # The failed booking left only the first appointment persisted.
    assert len(db.all()) == 1


def test_bad_datetime_raises(wired) -> None:
    with pytest.raises(AppointmentError):
        create_booking(
            name="X",
            department="Cardiology",
            doctor="Dr. Rajesh",
            preferred_date="next tuesday",
            preferred_time="afternoon",
        )


def _context(**overrides):
    from titan.services.tenant import BusinessContext

    base = {
        "business_id": 7,
        "business_name": "Sunrise Clinic",
        "slug": "sunrise",
        "timezone": "Asia/Kolkata",
        "business_hours": (10, 16),  # narrower than the env's 9-18
        "status": "active",
        "assistant_name": "Aarohi",
        "voice": "meera",
        "language": "en-IN",
    }
    base.update(overrides)
    return BusinessContext(**base)


def test_tenant_hours_reject_slot_env_hours_would_allow(wired, monkeypatch) -> None:
    # 09:30 is inside the env's 9-18 but outside the tenant's 10-16.
    monkeypatch.setattr(booking_mod, "get_business_context", lambda: _context())
    monkeypatch.setattr(booking_mod, "load_holidays", lambda business_id: set())

    with pytest.raises(SlotUnavailableError):
        create_booking(
            name="Early Bird",
            department="Cardiology",
            doctor="Dr. Rajesh",
            preferred_date="2026-07-08",
            preferred_time="09:30",
        )


def test_tenant_holiday_rejects_booking(wired, monkeypatch) -> None:
    from datetime import date

    monkeypatch.setattr(booking_mod, "get_business_context", lambda: _context())
    monkeypatch.setattr(
        booking_mod, "load_holidays", lambda business_id: {date(2026, 7, 8)}
    )

    with pytest.raises(SlotUnavailableError) as excinfo:
        create_booking(
            name="X",
            department="Cardiology",
            doctor="Dr. Rajesh",
            preferred_date="2026-07-08",
            preferred_time="12:00",
        )
    # A closed day offers no same-day alternatives.
    assert excinfo.value.alternatives == []


def test_tenant_booking_persists_business_id(wired, monkeypatch) -> None:
    settings, _cal, _dir, db, _report = wired
    monkeypatch.setattr(booking_mod, "get_business_context", lambda: _context())
    monkeypatch.setattr(booking_mod, "load_holidays", lambda business_id: set())

    # The appointments table enforces the businesses FK — seed the tenant row.
    from titan.db import connection as db_connection

    with db_connection.connect(settings.database_path) as conn:
        conn.execute(
            "INSERT INTO businesses (id, name, slug, timezone, created_at, updated_at)"
            " VALUES (7, 'Sunrise Clinic', 'sunrise', 'Asia/Kolkata',"
            " '2026-07-16T00:00:00', '2026-07-16T00:00:00')"
        )
        conn.commit()

    create_booking(
        name="Asha",
        department="Cardiology",
        doctor="Dr. Rajesh",
        preferred_date="2026-07-08",
        preferred_time="12:00",
    )
    rows = db.all()
    assert len(rows) == 1
    assert rows[0]["business_id"] == 7


def test_no_context_keeps_env_hours(wired, monkeypatch) -> None:
    monkeypatch.setattr(booking_mod, "get_business_context", lambda: None)
    booking = create_booking(
        name="Asha",
        department="Cardiology",
        doctor="Dr. Rajesh",
        preferred_date="2026-07-08",
        preferred_time="09:30",  # fine under the env's 9-18
    )
    assert booking.status == "confirmed"
